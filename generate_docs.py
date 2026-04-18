from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

def create_report():
    doc = Document()

    # --- STYLE ---
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # --- TITLE PAGE ---
    title = doc.add_heading('Project Report: Ayush-Guard', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Clinical Decision Support System (CDSS)')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('\n' * 5)
    
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run('DBMS Mini Project\n').bold = True
    info.add_run(f'Date: {datetime.date.today().strftime("%B %d, %Y")}\n')
    info.add_run('Institution: [Your Institution Here]\n')

    doc.add_page_break()

    # --- PROBLEM STATEMENT ---
    doc.add_heading('1. Problem Statement', level=1)
    doc.add_paragraph(
        "Ayush-Guard is a Clinical Decision Support System (CDSS) aimed at reconciling cross-disciplinary medication regimes "
        "(e.g., Allopathic and Ayurvedic medicines). The core problem it solves is adverse Drug-Drug Interactions (DDIs) "
        "and genetic sensitivity conflicts when a pharmacist attempts to dispense a new medicine without a holistic view "
        "of the patient's existing health profile. By integrating with the conceptual ABDM (Ayushman Bharat Digital Mission) "
        "framework, the platform extracts a patient's historical medical records and evaluates safety against accumulated health profiles."
    )

    # --- TOOLS & TECHNOLOGIES ---
    doc.add_heading('2. Tools and Technologies', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Component'
    hdr_cells[1].text = 'Technology Used'

    technologies = [
        ('Frontend', 'React, Vite, CSS'),
        ('Backend Architecture', 'Python, FastAPI, APIRouter'),
        ('Database', 'PostgreSQL (Supabase)'),
        ('Data Validation', 'Pydantic (Strict format enforcement)'),
        ('Testing', 'Pytest (Comprehensive route testing)'),
        ('Deployment', 'Vercel (Frontend), Render / Supabase (Backend & DB)')
    ]

    for component, tech in technologies:
        row_cells = table.add_row().cells
        row_cells[0].text = component
        row_cells[1].text = tech

    # --- DATABASE SCHEMA ---
    doc.add_heading('3. Database Schema', level=1)
    doc.add_paragraph(
        "We use Supabase (PostgreSQL) as our primary data layer. The system utilizes structured relational tables alongside "
        "flexible JSONB columns to mock the external FHIR-like records."
    )
    
    # Helper to add formatted bullets with optional indentations
    def add_bullet(text, bold_prefix="", indent=0):
        p = doc.add_paragraph(style='List Bullet')
        if indent > 0:
            p.paragraph_format.left_indent = Inches(0.5 * indent)
        if bold_prefix:
            p.add_run(bold_prefix).bold = True
        p.add_run(text)

    add_bullet(" Stores structural identity data (name, abha_id). No medical history is stored directly here to maintain separation of concerns.", "patients:")
    add_bullet(" A JSONB representation of the external ABDM gateway. It holds arrays of medication_history, pre_existing_conditions, allergies, and raw_fhir_bundles.", "abdm_mock_records:")
    add_bullet(" Connects patients to pharmacists, representing the digital consent framework (PENDING, ACCEPTED, REJECTED states).", "access_requests:")
    add_bullet(" Maps recursive patient IDs (e.g., Father to Son) to actively propagate genetic sensitivity warnings (e.g., G6PD Deficiency inheritance).", "family_relationships:")
    add_bullet(" Stores credentialing and license information for registered healthcare vendors.", "pharmacists & admins:")

    # --- SOLUTION IDEA ---
    doc.add_heading('4. Solution Idea & Architecture', level=1)
    doc.add_paragraph(
        "The project strongly separates concerns into a static UI layer (React/Vite) for highly responsive SPA mechanics, "
        "and a heavy computational API layer (FastAPI) for deterministic safety scoring."
    )
    
    steps = [
        "Data ETL Pipeline: final_generate.py ingests raw datasets (DrugBank, PharmGKB) into optimized, memory-resident JSON dictionaries (brand_to_salt, ddi_map, clinical_map).",
        "Step 1: Pharmacist issues an Access Request to a specific patient's ABHA ID.",
        "Step 2: Patient grants consent, unlocking their ABDM Mock Records.",
        "Step 3: Pharmacist inputs a brand name into the DashboardPanel.",
        "Step 4: Frontend calls POST /api/check-drug.",
        "Step 5: Backend resolves brand to salt and cross-references patient's active meds and genomics against the memory-resident JSON maps.",
        "Step 6: Backend checks family_relationships for inherited risks.",
        "Step 7: A structured alert (severity_tier, risk_probability, message) is returned to the user interface."
    ]
    
    for step in steps:
        doc.add_paragraph(step, style='List Number')

    # --- CHALLENGES FACED ---
    doc.add_heading('5. Challenges Faced', level=1)
    challenges = [
        ("Lack of Sensitive Clinical Data: ", "Due to privacy constraints, real patient datasets were unattainable. This occasionally caused the ML/NLP layers to hallucinate edge cases, forcing reliance on structured, mock demo injections to maintain deterministic safety scoring."),
        ("CORS Debugging: ", "Managing strict Cross-Origin Resource Sharing (CORS) rules between the isolated Vercel frontend and the Render-hosted backend during integration."),
        ("Data Parsing & Formats: ", "Handling semi-structured clinical JSON dumps required heavy usage of Pydantic to strictly enforce data validation across API boundaries."),
        ("Performance at Scale: ", "Fetching bulk user records strained network resources, requiring the implementation of backend pagination for efficient data retrieval in the GET /api/patients route."),
        ("Asynchronous Testing: ", "Ensuring that complex chained NLP endpoints evaluated safely under load necessitated rigorous Pytest implementations.")
    ]
    for title_text, desc in challenges:
        add_bullet(desc, title_text)

    # --- FUTURE IMPROVEMENTS ---
    doc.add_heading('6. Future Improvements', level=1)
    improvements = [
        ("Real ABDM Integration: ", "Swapping mocked JSON records for active HTTPS requests to the National Health Authority (NHA) gateway, decrypting FHIR JSONs via Diffie-Hellman keys."),
        ("Deep Learning Embeddings: ", "Upgrading text-matching to use a medical BERT model (e.g., ClinicalBERT) to identify semantic similarities within unstructured physician notes."),
        ("Advanced State Management: ", "Utilizing React Query or Redux to cache patient histories locally, drastically reducing redundant API trips."),
        ("Dosage-Specific Risk Calculation: ", "Upgrading the ML engine to require drug strength arrays, distinguishing between safe low doses and toxic thresholds.")
    ]
    for title_text, desc in improvements:
        add_bullet(desc, title_text)

    # --- CONCLUSION ---
    doc.add_heading('7. Conclusion', level=1)
    doc.add_paragraph(
        "Ayush-Guard successfully demonstrates the critical need for an interoperable Clinical Decision Support System. "
        "By merging modern web architecture (FastAPI/React) with robust deterministic mapping and mock ABDM integration, "
        "the prototype offers a scalable foundation for preventing cross-disciplinary drug interactions and protecting patient safety."
    )

    doc.add_page_break()

    # --- DELIVERABLES & RESOURCES ---
    doc.add_heading('Project Deliverables & Resources', level=1)
    doc.add_paragraph("[Drive Link: Demo Video] [ __________________________________________________ ]").bold = True
    doc.add_paragraph()
    doc.add_paragraph("[Drive Link: ER Diagram SQL & Image] [ __________________________________________________ ]").bold = True
    doc.add_paragraph()
    doc.add_paragraph("[Drive Link: Workflow Diagram] [ __________________________________________________ ]").bold = True
    
    doc.add_page_break()

    # --- SCRIPTS ---
    doc.add_heading('Scripts', level=1)
    
    doc.add_heading('ER Diagram (Mermaid Script)', level=2)
    mermaid_er = (
        "erDiagram\n"
        "    pharmacists ||--o{ patients : \"registers\"\n"
        "    pharmacists ||--o{ access_requests : \"makes\"\n"
        "    patients ||--o{ access_requests : \"receives\"\n"
        "    patients ||--o| abdm_mock_records : \"has\"\n"
        "    patients ||--o{ family_relationships : \"has relative\"\n\n"
        "    patients {\n"
        "        uuid id PK\n"
        "        text abha_id UK\n"
        "        text name\n"
        "        text phone\n"
        "        uuid registered_by FK\n"
        "    }\n"
        "    abdm_mock_records {\n"
        "        text abha_id PK, FK\n"
        "        jsonb basic_health_details\n"
        "        jsonb pre_existing_conditions\n"
        "        jsonb medication_history\n"
        "        jsonb allergies\n"
        "    }\n"
        "    access_requests {\n"
        "        uuid id PK\n"
        "        uuid pharmacist_id FK\n"
        "        uuid patient_id FK\n"
        "        text status\n"
        "    }\n"
        "    family_relationships {\n"
        "        uuid id PK\n"
        "        uuid patient_id FK\n"
        "        uuid relative_id FK\n"
        "        text relationship_type\n"
        "    }\n"
        "    pharmacists {\n"
        "        uuid id PK\n"
        "        text name\n"
        "        text license_number UK\n"
        "    }\n"
        "    admins {\n"
        "        uuid id PK\n"
        "        text username UK\n"
        "    }"
    )
    p_er = doc.add_paragraph(mermaid_er)
    for run in p_er.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)

    doc.add_heading('Workflow Diagram (Mermaid Script)', level=2)
    mermaid_wf = (
        "flowchart TD\n"
        "    Start([Pharmacist wants to dispense drug]) --> Step1[Request Access to Patient Profile]\n"
        "    Step1 --> Step2{Consent Granted?}\n"
        "    Step2 -- Yes --> Step3[Fetch ABDM Records JSON]\n"
        "    Step2 -- No --> End1([Access Denied])\n"
        "    Step3 --> Step4[Input Brand: POST /api/check-drug]\n"
        "    Step4 --> Step5[Resolve Brand to Salt via final_generate maps]\n"
        "    Step5 --> Step6[Cross-reference DDI, Genotype & Allergy Maps]\n"
        "    Step6 --> Step7[Check family_relationships for Inheritance Risks]\n"
        "    Step7 --> Step8[Calculate Safety Score & Output Alerts]\n"
        "    Step8 --> End([Dispense or Deny Medication])\n\n"
        "    classDef step fill:#f9f9f9,stroke:#333,stroke-width:2px;\n"
        "    class Step1,Step2,Step3,Step4,Step5,Step6,Step7,Step8 step;"
    )
    p_wf = doc.add_paragraph(mermaid_wf)
    for run in p_wf.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)

    # --- FOOTER ---
    section = doc.sections[0]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.text = "Generated by Ayush-Guard CDSS - Internal Documentation"
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save('Project_Report_AyushGuard.docx')
    print("Report generated successfully: Project_Report_AyushGuard.docx")

if __name__ == "__main__":
    create_report()